"""
Text extraction module for PDF, DOCX, and TXT files.
"""
import os
from pathlib import Path
from typing import Optional

try:
    from PyPDF2 import PdfReader
except ImportError:
    from pypdf import PdfReader

from docx import Document as DocxDocument


class TextExtractor:
    """Extract text from various document formats"""

    @staticmethod
    def extract_from_pdf(file_path: str) -> str:
        """
        Extract text from PDF file

        Args:
            file_path: Path to the PDF file

        Returns:
            Extracted text content
        """
        try:
            reader = PdfReader(file_path)
            text_parts = []

            for page in reader.pages:
                try:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
                except Exception:
                    continue

            return '\n'.join(text_parts)
        except Exception as e:
            return f"Error extracting PDF: {str(e)}"

    @staticmethod
    def extract_from_docx(file_path: str) -> str:
        """
        Extract text from Word DOCX file

        Args:
            file_path: Path to the DOCX file

        Returns:
            Extracted text content
        """
        try:
            doc = DocxDocument(file_path)
            text_parts = []

            # Extract from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text_parts.append(paragraph.text)

            # Extract from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            text_parts.append(cell.text)

            return '\n'.join(text_parts)
        except Exception as e:
            return f"Error extracting DOCX: {str(e)}"

    @staticmethod
    def extract_from_txt(file_path: str, encoding: str = 'utf-8') -> str:
        """
        Extract text from TXT file

        Args:
            file_path: Path to the TXT file
            encoding: File encoding (default: utf-8)

        Returns:
            File content
        """
        try:
            # Try multiple encodings
            encodings = [encoding, 'utf-8', 'latin-1', 'cp1252', 'iso-8859-1']

            for enc in encodings:
                try:
                    with open(file_path, 'r', encoding=enc) as f:
                        return f.read()
                except UnicodeDecodeError:
                    continue

            return "Error: Could not decode file with any common encoding"
        except Exception as e:
            return f"Error extracting TXT: {str(e)}"

    @classmethod
    def extract_text(cls, file_path: str, file_type: str) -> str:
        """
        Extract text based on file type

        Args:
            file_path: Path to the file
            file_type: Type of file (pdf, docx, txt)

        Returns:
            Extracted text content
        """
        file_type = file_type.lower()

        if file_type == 'pdf':
            return cls.extract_from_pdf(file_path)
        elif file_type == 'docx':
            return cls.extract_from_docx(file_path)
        elif file_type == 'txt':
            return cls.extract_from_txt(file_path)
        else:
            return f"Unsupported file type: {file_type}"

    @staticmethod
    def clean_text(text: str) -> str:
        """
        Clean extracted text by removing excessive whitespace

        Args:
            text: Raw text

        Returns:
            Cleaned text
        """
        if not text:
            return ""

        # Remove extra whitespace and newlines
        lines = text.split('\n')
        cleaned_lines = []

        for line in lines:
            line = line.strip()
            if line:
                cleaned_lines.append(line)

        return '\n'.join(cleaned_lines)

    @staticmethod
    def get_word_count(text: str) -> int:
        """Get word count of text"""
        if not text:
            return 0
        return len(text.split())

    @staticmethod
    def get_character_count(text: str) -> int:
        """Get character count of text"""
        return len(text) if text else 0
